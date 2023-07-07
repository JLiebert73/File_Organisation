import os
from PIL import Image
import PyPDF2
import docx
import openpyxl
import pptx

def extract_image_metadata(image_path):
    with Image.open(image_path) as img:
        # Extract specific metadata properties
        # Example: Extract image format, size, and mode
        format = img.format
        size = img.size
        mode = img.mode

        # Return metadata as desired
        return format, size, mode

def extract_pdf_metadata(pdf_path):
    with open(pdf_path, 'rb') as file:
        pdf = PyPDF2.PdfFileReader(file)
        metadata = pdf.getDocumentInfo()

        # Extract specific metadata properties
        # Example: Extract number of pages, title, author, and subject
        num_pages = pdf.getNumPages()
        title = metadata.title
        author = metadata.author
        subject = metadata.subject

        # Return metadata as desired
        return num_pages, title, author, subject

def extract_docx_metadata(docx_path):
    doc = docx.Document(docx_path)
    properties = doc.core_properties

    # Extract specific metadata properties
    # Example: Extract title, author, and created date
    title = properties.title
    author = properties.author
    created = properties.created

    # Return metadata as desired
    return title, author, created

def extract_xlsx_metadata(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path)
    properties = workbook.properties

    # Extract specific metadata properties
    # Example: Extract title, author, and created date
    title = properties.title
    author = properties.creator
    created = properties.created

    # Return metadata as desired
    return title, author, created

def extract_pptx_metadata(pptx_path):
    prs = pptx.Presentation(pptx_path)
    core_props = prs.core_properties

    # Extract specific metadata properties
    # Example: Extract title, author, and created date
    title = core_props.title
    author = core_props.author
    created = core_props.created

    # Return metadata as desired
    return title, author, created

def compile_metadata(directory):
    metadata = []
    
    for filename in os.listdir(directory):
        file_path = os.path.join(directory, filename)
        
        if os.path.isfile(file_path):
            # Extract metadata based on file type
            file_extension = os.path.splitext(filename)[1].lower()
            
            if file_extension in ['.jpg', '.jpeg', '.png', '.gif']:
                metadata.append(('Image', file_path, *extract_image_metadata(file_path)))
                
            elif file_extension == '.pdf':
                metadata.append(('PDF', file_path, *extract_pdf_metadata(file_path)))
                
            elif file_extension == '.docx':
                metadata.append(('DOCX', file_path, *extract_docx_metadata(file_path)))
                
            elif file_extension == '.xlsx':
                metadata.append(('XLSX', file_path, *extract_xlsx_metadata(file_path)))
                
            elif file_extension == '.pptx':
                metadata.append(('PPTX', file_path, *extract_pptx_metadata(file_path)))
                
            # Add conditions for other file types as needed
        
    return metadata

# Example usage:
path = os.getcwd() 
directory = r"C:\Users\ASUS\Downloads\Text"
metadata = compile_metadata(directory)

# Print the compiled metadata
for item in metadata:
    print('File:', item[1])
    print('Type:', item[0])
    print('Metadata:', item[2:])
    print()
