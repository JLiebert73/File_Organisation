import os
from PIL import Image
import PyPDF2
import docx
import openpyxl
import pptx
from tabulate import tabulate


def extract_image_metadata(image_path):
    with Image.open(image_path) as img:
        format = img.format
        size = img.size
        mode = img.mode
        return format, size, mode


def extract_pdf_metadata(pdf_path):
    with open(pdf_path, 'rb') as file:
        pdf = PyPDF2.PdfReader(file)
        metadata = pdf.metadata
        num_pages = len(pdf.pages)
        title = metadata.get('/Title')
        author = metadata.get('/Author')
        subject = metadata.get('/Subject')
        return num_pages, title, author, subject


def extract_docx_metadata(docx_path):
    doc = docx.Document(docx_path)
    properties = doc.core_properties
    title = properties.title
    author = properties.author
    created = properties.created
    return title, author, created


def extract_xlsx_metadata(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path)
    properties = workbook.properties
    title = properties.title
    author = properties.creator
    created = properties.created
    return title, author, created


def extract_pptx_metadata(pptx_path):
    prs = pptx.Presentation(pptx_path)
    core_props = prs.core_properties
    title = core_props.title
    author = core_props.author
    created = core_props.created
    return title, author, created


def extract_txt_metadata(txt_path):
    with open(txt_path, 'r') as file:
        content = file.read()
        return content.strip()


def compile_metadata(directory, file_extension):
    metadata = []

    for filename in os.listdir(directory):
        file_path = os.path.join(directory, filename)

        if os.path.isfile(file_path):
            if os.path.splitext(filename)[1].lower() == file_extension:
                if file_extension in ['.jpg', '.jpeg', '.png', '.gif']:
                    metadata.append(('Image', file_path, *extract_image_metadata(file_path)))

                elif file_extension == '.pdf':
                    metadata.append(('PDF', file_path, *extract_pdf_metadata(file_path)))

                elif file_extension == '.docx':
                    metadata.append(('DOCX', file_path, *extract_docx_metadata(file_path)))

                elif file_extension == '.xlsx':
                    metadata.append(('XLSX', file_path, *extract_xlsx_metadata(file_path)))

                elif file_extension == '.pptx':
                    metadata.append(('PPTX', file_path, *extract_pptx_metadata(file_path)))

                elif file_extension == '.png':
                    metadata.append(('PNG', file_path, *extract_image_metadata(file_path)))

                elif file_extension == '.doc':
                    metadata.append(('DOC', file_path, *extract_docx_metadata(file_path)))

                elif file_extension == '.txt':
                    metadata.append(('TXT', file_path, extract_txt_metadata(file_path)))

    return metadata

